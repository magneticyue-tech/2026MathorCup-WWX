"""Read-only audit of the eight MathorCup attachments.

No raw file is modified. Numeric parsing is used only for diagnostics;
all original CSV strings and XLSX coordinates remain available as evidence.
Usage: python src/audit_data.py [--root PATH]
Dependencies: Python 3.11+, numpy, pandas, openpyxl.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import platform
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

MISSING = {"", "NA", "N/A", "NaN", "nan", "NULL", "null", "None", "#N/A"}


def save_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str, allow_nan=False), encoding="utf-8")


def scalar(x):
    if x is None or pd.isna(x):
        return None
    if isinstance(x, np.generic):
        return x.item()
    return x


def write_csv(path, rows):
    if not rows:
        return
    pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8-sig")


def read_csv_raw(path):
    content = path.read_bytes()
    encoding = next(enc for enc in ("utf-8-sig", "utf-16", "cp1252") if decodes(content, enc))
    text = content.decode(encoding)
    dialect = csv.Sniffer().sniff("\n".join(text.splitlines()[:5]), delimiters=",;\t")
    rows = list(csv.reader(text.splitlines(), dialect))
    widths = Counter(map(len, rows))
    if len(widths) != 1:
        raise ValueError(f"Inconsistent CSV widths in {path}: {widths}")
    header = rows[0]
    if len(set(header)) != len(header):
        raise ValueError(f"Duplicate CSV headers in {path}")
    return pd.DataFrame(rows[1:], columns=header), {"encoding": encoding, "delimiter": dialect.delimiter, "row_widths": dict(widths)}


def decodes(content, encoding):
    try:
        content.decode(encoding)
        return True
    except UnicodeError:
        return False


def csv_audit(path, out):
    raw, parsing = read_csv_raw(path)
    missing = raw.apply(lambda col: col.map(lambda x: x.strip() in MISSING))
    numeric = raw.mask(missing).apply(pd.to_numeric, errors="coerce")
    rows = []
    for j, name in enumerate(raw.columns):
        s = raw[name]
        n = numeric[name].dropna()
        attempted = (~missing[name]).sum()
        full_numeric = bool(len(n) == attempted and attempted)
        row = {"file": path.name, "position": j + 1, "column": name,
               "nonmissing": int((~missing[name]).sum()), "missing": int(missing[name].sum()),
               "missing_rate": float(missing[name].mean()), "unique_nonmissing": int(s[~missing[name]].nunique()),
               "numeric_count": int(len(n)), "fully_numeric": full_numeric,
               "missing_tokens": json.dumps(dict(Counter(s[missing[name]])), ensure_ascii=False),
               "examples": json.dumps(s[~missing[name]].drop_duplicates().head(5).tolist(), ensure_ascii=False)}
        for label, val in (("min", n.min()), ("q01", n.quantile(.01)), ("q25", n.quantile(.25)), ("median", n.median()),
                           ("q75", n.quantile(.75)), ("q99", n.quantile(.99)), ("max", n.max()), ("mean", n.mean()), ("std_ddof1", n.std())):
            row[label] = scalar(val)
        row.update(negative=int((n < 0).sum()), zero=int((n == 0).sum()), positive=int((n > 0).sum()),
                   integer_count=int(np.isclose(n, np.round(n), atol=1e-10, rtol=0).sum()),
                   nonfinite=int((~np.isfinite(n)).sum()))
        if full_numeric and len(n):
            q1, q3 = n.quantile([.25, .75]); iqr = q3 - q1
            row["tukey_1_5_iqr_flags"] = int(((n < q1 - 1.5*iqr) | (n > q3 + 1.5*iqr)).sum())
        rows.append(row)
    key_cols = [x for x in raw.columns if any(k in x.lower() for k in ("beer", "name", "style", "category", "id"))]
    row_rows = [{"csv_row": i+2, "missing": int(missing.iloc[i].sum()),
                 **{c: raw.iloc[i][c] for c in key_cols[:8]}} for i in range(len(raw))]
    write_csv(out / f"{path.stem}_columns.csv", rows)
    write_csv(out / f"{path.stem}_rows.csv", row_rows)
    missing_cells = [{"csv_row": int(i)+2, "column": raw.columns[j], "token": raw.iat[i,j]}
                     for i,j in zip(*np.where(missing.to_numpy()))]
    write_csv(out / f"{path.stem}_missing_cells.csv", missing_cells)
    info = {"file": path.name, "rows": len(raw), "columns": len(raw.columns), "parsing": parsing,
            "headers": list(raw.columns), "missing_cells": int(missing.to_numpy().sum()),
            "missing_columns": {c: int(missing[c].sum()) for c in raw.columns if missing[c].any()},
            "missing_row_count": int(missing.any(axis=1).sum()),
            "duplicate_full_rows": int(raw.duplicated().sum()),
            "constant_columns": [r["column"] for r in rows if r["unique_nonmissing"] <= 1],
            "column_profiles": rows,
            "head": raw.iloc[:3,:min(12,len(raw.columns))].to_dict(orient="records")}
    return info, raw, numeric


def xlsx_audit(path, out):
    wb = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    info = {"file": path.name, "properties": {k: str(getattr(wb.properties,k)) for k in ("creator","title","description","created","modified")},
            "sheets": [], "defined_names": list(wb.defined_names)}
    cells = []
    for ws in wb:
        occupied = [c for row in ws for c in row if c.value is not None]
        rr = [c.row for c in occupied]; cc = [c.column for c in occupied]
        texts = [{"cell": c.coordinate, "value": str(c.value)} for c in occupied if isinstance(c.value,str)]
        si = {"name": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column,
              "nonempty_cells": len(occupied), "occupied_bounds": [min(rr),min(cc),max(rr),max(cc)] if occupied else None,
              "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
              "formula_count": sum(c.data_type == 'f' for c in occupied), "text_cells": texts,
              "top_rows": [[scalar(c.value) for c in row] for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,12),max_col=min(ws.max_column,14))]}
        for c in occupied:
            cells.append({"file": path.name, "sheet": ws.title, "cell": c.coordinate, "row": c.row, "column": c.column,
                          "type": c.data_type, "value": c.value, "cached": cached[ws.title][c.coordinate].value,
                          "number_format": c.number_format, "comment": c.comment.text if c.comment else ""})
        info["sheets"].append(si)
    write_csv(out / f"{path.stem}_xlsx_cells.csv", cells)
    save_json(out / f"{path.stem}_xlsx_structure.json", info)
    return info


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args(); root = args.root.resolve()
    rawdir = root / "data" / "raw"; out = root / "results" / "data_audit"; out.mkdir(parents=True, exist_ok=True)
    manifest = [{"file": p.name, "bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                for p in sorted(rawdir.iterdir()) if p.is_file() and p.name != '.gitkeep']
    save_json(out / "raw_manifest.json", manifest)
    tables = {}; csvs = {}; nums = {}; columns = []
    for p in sorted(rawdir.glob("*.csv")):
        info, raw, num = csv_audit(p,out); tables[p.stem] = info; csvs[p.stem] = raw; nums[p.stem] = num
        columns.extend(info["column_profiles"])
    write_csv(out / "variable_dictionary.csv", columns)
    xlsxs = {p.stem: xlsx_audit(p,out) for p in sorted(rawdir.glob("*.xlsx"))}
    save_json(out / "audit_summary.json", {"versions": {"python":platform.python_version(), "pandas":pd.__version__,"numpy":np.__version__,"openpyxl":openpyxl.__version__},"csv": tables,"xlsx": xlsxs})
    compact = {k: {x:v[x] for x in ("rows","columns","headers","missing_cells","missing_columns","missing_row_count","duplicate_full_rows","constant_columns","head")} for k,v in tables.items()}
    compact["xlsx"] = {k:[{x:s[x] for x in ("name","max_row","max_column","nonempty_cells","formula_count","top_rows")} for s in v["sheets"]] for k,v in xlsxs.items()}
    save_json(out / "initial_overview.json", compact)
    print(json.dumps({k:{x:v[x] for x in ("rows","columns","missing_cells","missing_columns")} for k,v in tables.items()},ensure_ascii=False,indent=2))
    print("XLSX sheets", {k:[(s['name'],s['max_row'],s['max_column']) for s in v['sheets']] for k,v in xlsxs.items()})


if __name__ == "__main__":
    main()
